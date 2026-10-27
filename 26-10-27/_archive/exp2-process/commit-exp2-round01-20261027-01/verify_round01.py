#!/usr/bin/env python3
"""Read-only integrity checks for the experiment-2 round-01 delivery."""
import argparse
import hashlib
import json
import re
import statistics
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook
from PIL import Image


def sha(data):
    return hashlib.sha256(data).hexdigest()


def check(test, message):
    if not test:
        raise AssertionError(message)


def score(reference, response):
    expected = list(re.sub(r'\s+', '', reference))
    predicted = []
    if response.strip() != '[UNREADABLE]':
        for part in re.split(r'(\[UNK\])', response):
            predicted.extend([None] if part == '[UNK]' else list(re.sub(r'\s+', '', part)))
    previous = [(j, 0, 0, 0, j) for j in range(len(predicted) + 1)]
    for i, a in enumerate(expected, 1):
        current = [(i, 0, 0, i, 0)]
        for j, b in enumerate(predicted, 1):
            c, m, s, d, ins = previous[j - 1]
            same = a == b
            candidates = [(c + int(not same), m - int(same), s + int(not same), d, ins)]
            c, m, s, d, ins = previous[j]
            candidates.append((c + 1, m, s, d + 1, ins))
            c, m, s, d, ins = current[-1]
            candidates.append((c + 1, m, s, d, ins + 1))
            current.append(min(candidates))
        previous = current
    edits, negative_matches, substitutions, deletions, insertions = previous[-1]
    return {
        'correct': -negative_matches, 'substitutions': substitutions,
        'deletions': deletions, 'insertions': insertions, 'edit_distance': edits,
        'total': len(expected), 'accuracy': -negative_matches / len(expected),
        'cer': edits / len(expected), 'unk_tokens': response.count('[UNK]'),
        'unreadable': response.strip() == '[UNREADABLE]',
        'exact_transcription': response.strip('\r\n') == reference.strip('\r\n'),
    }


def approx(a, b):
    return isinstance(a, (int, float)) and abs(a - b) < 1e-9


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('phase', choices=['baseline', 'modified', 'rollback', 'checkout'])
    parser.add_argument('--repo', type=Path, required=True)
    parser.add_argument('--base', type=Path, required=True)
    parser.add_argument('--archives', type=Path)
    args = parser.parse_args()
    phase, repo, base = args.phase, args.repo.resolve(), args.base.resolve()
    evidence = base / 'evidence-01'
    records = evidence / 'records'
    reference = (repo / '26-10-26/Random-characters.txt').read_text(encoding='utf-8')
    image_rows = json.loads((records / 'image-manifest.json').read_text(encoding='utf-8'))
    image_index = {r['case_id']: r for r in image_rows}
    record_paths = sorted(records.glob('exp2-result-*.json'))
    cases = {f'f{font}-g{gray:02}' for font in (16, 24, 32, 40, 64) for gray in range(1, 33)}
    check(len(record_paths) == len(image_rows) == len(image_index) == 160, 'record or manifest count')
    workbook = load_workbook(base / '实验结果-01.xlsx', read_only=True, data_only=True)
    formula_book = load_workbook(base / '实验结果-01.xlsx', read_only=True, data_only=False)
    check(workbook.sheetnames == ['实验结果', '完整回复'], 'workbook sheet names')
    sheet, replies = workbook.worksheets
    check(formula_book['实验结果']['H4'].value == '=COUNTIFS(B10:B169,100,F10:F169,0)', 'full-score formula')
    result_rows = {}
    for i in range(10, 170):
        shown = sheet.cell(i, 1).value
        check(isinstance(shown, str) and re.fullmatch(r'(?:f|px)(16|24|32|40|64)-g(?:0[1-9]|[12][0-9]|3[0-2])', shown), f'workbook condition row {i}')
        case = 'f' + shown[2:] if shown.startswith('px') else shown
        check(case not in result_rows, f'duplicate workbook condition {case}')
        check(replies.cell(i - 1, 1).value == shown, f'workbook ID mismatch {case}')
        result_rows[case] = i
    check(set(result_rows) == cases, 'workbook condition set')
    archive_cache = {}
    current = {}
    total = 0
    for path in record_paths:
        d = json.loads(path.read_text(encoding='utf-8'))
        case = d['case_id']
        check(case in cases and case not in current, f'record ID {case}')
        check(d['status'] == 'completed' and d['round'] == d['repeat'] == 1, f'record status {case}')
        check(case == f"f{d['font_px']}-g{d['delta_g']:02}", f'record parameters {case}')
        check(d['screen_brightness_mode'] == 0 and d['screen_brightness'] == 255 and
              d['brightness_after'] == {'mode': 0, 'value': 255}, f'brightness {case}')
        m = image_index[case]
        check(m['font_px'] == d['font_px'] and m['delta_g'] == d['delta_g'] and
              m['expected_text'] == reference and m['png_sha256'] == d['source_sha256'] == d['phone_source_sha256'], f'image manifest {case}')
        check(d['reference_sha256'] == sha(reference.encode()), f'reference hash {case}')
        paths = [('source-images/' + d['image'], 'source_sha256'), (d['response_file'], 'response_sha256'),
                 (d['screenshot'], 'screenshot_sha256'), (d['source_screen'], 'source_screen_sha256')]
        for rel, key in paths:
            p = evidence / rel
            check(p.is_file() and sha(p.read_bytes()) == d[key], f'asset {case} {key}')
        text = (evidence / d['response_file']).read_text(encoding='utf-8').rstrip('\r\n')
        check(text == d['response'].rstrip('\r\n'), f'response {case}')
        recalculated = score(reference, text)
        for key, value in recalculated.items():
            check(d.get(key) == value, f'rescore {case} {key}')
        total += d['correct']
        with Image.open(evidence / 'source-images' / d['image']) as source, Image.open(evidence / d['source_screen']) as screen:
            check(source.size == screen.size == (1080, 2400), f'input dimensions {case}')
            bounds = (65, 650, 1030, 1120)
            check(source.crop(bounds).convert('RGB').tobytes() == screen.crop(bounds).convert('RGB').tobytes(), f'source pixels {case}')
        with Image.open(evidence / d['screenshot']) as screenshot:
            check(screenshot.size == (1080, 2400), f'screenshot dimensions {case}')
            screenshot.verify()
        row = result_rows[case]
        numbers = [d['correct'], d['accuracy'], d['substitutions'], d['deletions'], d['insertions'], d['unk_tokens']]
        for col, value in enumerate(numbers, 2):
            check(approx(sheet.cell(row, col).value, value), f'workbook score {case} {col}')
        for col, value in zip((10, 11, 12, 13), (d['font_px'], d['delta_g'], 1, d['cer'])):
            check(approx(sheet.cell(row, col).value, value), f'workbook parameters {case} {col}')
        check((replies.cell(row - 1, 2).value or '').rstrip('\r\n') == text and
              replies.cell(row - 1, 3).value == Path(d['screenshot']).name, f'workbook response {case}')
        if phase in ('baseline', 'rollback'):
            check(args.archives is not None, 'archive directory required')
            name = Path(d['ui_archive']).name
            if name not in archive_cache:
                z = zipfile.ZipFile(args.archives / name)
                archive_cache[name] = (z, json.loads(z.read('manifest.json'))['files'])
            z, manifest = archive_cache[name]
            answer = z.read(d['answer_xml_archive_member'])
            prompt = z.read(d['prompt_xml_archive_member'])
            check(sha(answer) == manifest[d['answer_xml_archive_member']] and sha(prompt) == manifest[d['prompt_xml_archive_member']], f'archive XML hash {case}')
        else:
            answer = (evidence / 'ui' / case / 'answer.xml').read_bytes()
            prompt = (evidence / 'ui' / case / 'prompt.xml').read_bytes()
        ax, px = ET.fromstring(answer), ET.fromstring(prompt)
        extracted = '\n'.join(x.get('text') for x in ax.iter('node') if x.get('class') == 'android.widget.TextView' and x.get('text'))
        edits = [x for x in px.iter('node') if x.get('class') == 'android.widget.EditText']
        check(extracted == text, f'answer XML {case}')
        check(len(edits) == 1 and edits[0].get('text') == d['prompt_sent'], f'prompt XML {case}')
        current[case] = d
    for z, _ in archive_cache.values():
        z.close()
    check(set(current) == cases and total == 5724, 'total or record set')
    check(approx(sheet['B4'].value, 5724 / 16000) and sheet['F4'].value == 160 and sheet['H4'].value == 0, 'workbook summary')
    historical = json.loads((records / 'baseline48-reference.json').read_text(encoding='utf-8'))['rows']
    check(len(historical) == 96, 'historical count')
    for i, h in enumerate(historical, 9):
        p = repo / h['record_path']
        check(p.is_file() and sha(p.read_bytes()) == h['record_sha256'], f'historical hash {h["record_path"]}')
        wanted = [h['round'], h['delta_g'], 48, h['correct'], h['accuracy'], h['substitutions'], h['deletions'], h['insertions'], h['cer'], h['record_path']]
        check([replies.cell(i, c).value for c in range(5, 15)] == wanted, f'historical row {i}')
    for gray in range(1, 33):
        for col, font in ((16, 16), (17, 24), (18, 32), (19, 40), (21, 64)):
            check(approx(sheet.cell(9 + gray, col).value, current[f'f{font}-g{gray:02}']['accuracy']), f'matrix {font}/{gray}')
        hist_mean = statistics.mean(h['accuracy'] for h in historical if h['delta_g'] == gray)
        check(approx(sheet.cell(9 + gray, 20).value, hist_mean), f'historical matrix {gray}')
    replacements = json.loads((records / 'quality-retest-20261027.json').read_text(encoding='utf-8'))['cases']
    check(len(replacements) == 4, 'retest count')
    for r in replacements:
        case = r['case_id']
        old = base / r['old_evidence_dir'] / 'records' / f'exp2-result-{case}.json'
        new = records / f'exp2-result-{case}.json'
        check(sha(old.read_bytes()) == r['old_record_sha256'] and sha(new.read_bytes()) == r['new_record_sha256'], f'retest provenance {case}')
        gate = current[case]['source_fidelity']
        check(gate['status'] == 'PASS' and gate['consecutive_exact'] == 2 and all(x['exact'] for x in gate['observations']), f'retest gate {case}')
    indexed = 0
    if phase in ('modified', 'checkout'):
        index = json.loads((evidence / 'UI证据索引.json').read_text(encoding='utf-8'))
        manifests = {}
        for p in (evidence / 'ui/archive-manifests').glob('*.manifest.json'):
            name = p.name.removesuffix('.manifest.json')
            manifests[name] = json.loads(p.read_text(encoding='utf-8'))['files']
        check(len(manifests) == 2, 'source archive manifests')
        seen = set()
        for item in index['entries']:
            rel = item['file']
            check(rel not in seen, f'duplicate evidence index {rel}')
            seen.add(rel)
            p = evidence / rel
            check(p.is_file() and sha(p.read_bytes()) == item['sha256'], f'evidence index hash {rel}')
            if item.get('source_archive') and item['source_member'] != 'manifest.json':
                check(manifests[item['source_archive']].get(item['source_member']) == item['sha256'], f'archive manifest provenance {rel}')
        indexed = len(seen)
        for case in cases:
            for member in ('answer.xml', 'input.xml', 'prompt.xml', 'source-clean.xml', 'source-initial.xml', 'source-query.txt', 'state.json'):
                check(f'ui/{case}/{member}' in seen, f'missing indexed evidence {case} {member}')
        for r in replacements:
            case = r['case_id']
            for member in ('answer.xml', 'input.xml', 'prompt.xml', 'source-clean.xml', 'source-initial.xml', 'source-query.txt', 'state.json'):
                check(f'excluded/fidelity-superseded-20261027/{case}/ui/{member}' in seen, f'missing superseded evidence {case} {member}')
            for member in ('source-fidelity.json', 'source-fidelity/01.png', 'source-fidelity/02.png'):
                check(f'ui/{case}/{member}' in seen, f'missing retest gate evidence {case} {member}')
            gate = current[case]['source_fidelity']
            check(json.loads((evidence / 'ui' / case / 'source-fidelity.json').read_text(encoding='utf-8')) == gate, f'retest gate JSON {case}')
            with Image.open(evidence / 'source-images' / current[case]['image']) as source:
                expected = source.crop(tuple(gate['roi'])).convert('RGB').tobytes()
            check(sha(expected) == gate['expected_roi_sha256'], f'retest gate reference {case}')
            for obs in gate['observations']:
                frame = evidence / 'ui' / case / 'source-fidelity' / f'{obs["index"]:02}.png'
                check(sha(frame.read_bytes()) == obs['source_screen_sha256'], f'retest gate frame hash {case}')
                with Image.open(frame) as image:
                    actual = image.crop(tuple(gate['roi'])).convert('RGB').tobytes()
                check(actual == expected, f'retest gate frame pixels {case}')
        notes = (base / '实验说明.md').read_text(encoding='utf-8')
        for target in re.findall(r'\[[^\]]+\]\(([^)]+)\)', notes):
            if target.startswith(('https:', 'http:', '#')):
                continue
            relative = unquote(target.split('#', 1)[0])
            dest = (base / relative).resolve()
            if not dest.exists() and relative.startswith('../../26-10-26/'):
                # A detached test copy is not at the same depth as its Git checkout.
                dest = (repo / relative[len('../../'):]).resolve()
            check(dest.exists(), f'broken documentation link {target}')
        check('px16-g01' in notes and 'f16-g01' in notes, 'workbook alias documentation')
    print(f'{phase.upper()}: cases=160 asset_hashes=640 source_pixels=160/160')
    print(f'{phase.upper()}: scores=5724/16000 answer_xml=160 prompt_xml=160 historical=96 retests=4')
    print(f'{phase.upper()}: workbook_rows=160 excel_ids=px-to-f indexed_files={indexed}')
    print(f'{phase.upper()}: PASS')


if __name__ == '__main__':
    main()
